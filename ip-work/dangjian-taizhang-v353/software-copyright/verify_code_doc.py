#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""校验生成的《源程序》Word 文档是否符合软著要求"""
import sys
import unicodedata
from docx import Document
from docx.oxml.ns import qn

DOC = '/workspace/ip-work/dangjian-taizhang-v353/software-copyright/delivery/党建工作台账_源程序.docx'
LIMIT = 95  # 允许的最大显示宽度


def disp_width(s):
    return sum(2 if unicodedata.east_asian_width(c) in ('W', 'F') else 1 for c in s)


def main():
    doc = Document(DOC)
    body = doc.element.body
    paras = doc.paragraphs

    # 仅统计正文段落（非空页眉/页脚不计入 doc.paragraphs）
    print(f'正文段落数: {len(paras)}')

    # 分页符统计
    breaks = 0
    for p in paras:
        if p.paragraph_format.page_break_before:
            breaks += 1
    print(f'page_break_before 数: {breaks}')

    # 节数
    print(f'节数: {len(doc.sections)}')
    for i, s in enumerate(doc.sections):
        hdr = ' | '.join(x.text for x in s.header.paragraphs if x.text.strip())
        print(f'  节{i + 1}: 页眉="{hdr}"  页面={s.page_width.cm:.1f}x{s.page_height.cm:.1f}cm '
              f'边距(上/下/左/右)={s.top_margin.cm:.1f}/{s.bottom_margin.cm:.1f}/'
              f'{s.left_margin.cm:.1f}/{s.right_margin.cm:.1f}cm')

    # 超宽行检查
    over = [(i, disp_width(p.text), p.text[:60]) for i, p in enumerate(paras) if disp_width(p.text) > LIMIT]
    print(f'超过 {LIMIT} 显示宽度的行数: {len(over)}')
    for i, w, t in over[:5]:
        print(f'  - 段落{i} 宽度{w}: {t}')

    # 行距与字号抽查
    p0 = paras[0]
    pf = p0.paragraph_format
    print(f'首段行距规则={pf.line_spacing_rule} 行距={pf.line_spacing}')
    print(f'首段字号={p0.runs[0].font.size.pt if p0.runs and p0.runs[0].font.size else None}')

    # 每页行数校验：按分页符切分
    pages, cur = [], 0
    for p in paras:
        if p.paragraph_format.page_break_before and cur > 0:
            pages.append(cur)
            cur = 0
        cur += 1
    pages.append(cur)
    print(f'按分页符推算页数: {len(pages)}')
    bad = [(i + 1, n) for i, n in enumerate(pages) if n != 50]
    print(f'非 50 行的页: {bad if bad else "无（全部为 50 行）"}')
    print(f'每部分页数: 前 {pages[:30] and len(pages[:30])} 页 / 后 {len(pages[30:])} 页')


if __name__ == '__main__':
    main()